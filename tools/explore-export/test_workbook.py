import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from openpyxl import load_workbook
from workbook import create


class Workbook(unittest.TestCase):
    def data(self):
        return dict(title='Test export',when='tonight',exported='2026-10-30T17:00:00Z',
                    windowStart='2026-10-30T23:00:00Z',windowEnd='2026-10-31T08:00:00Z',runId='test',
                    warnings=[],records=[],venues=[],tables=[dict(name='Listings',
                    headers=['Listing #','Map pin','Category','Title','My research'],notes=['My research'],
                    rows=[[1,1,'Events','=IMPORTXML("https://example.org", "//x")','']])])

    def test_untrusted_strings_stay_literal_and_research_is_editable(self):
        with TemporaryDirectory() as temp:
            path=Path(temp)/'listings.xlsx';create(self.data(),path)
            book=load_workbook(path)
            self.assertEqual(book['Listings']['D2'].data_type,'s')
            self.assertTrue(book['Listings']['D2'].value.startswith('='))
            self.assertIsNone(book['Listings']['E2'].value)
            self.assertEqual(book['Listings'].freeze_panes,'E2')
            self.assertEqual(book['Listings'].auto_filter.ref,'A1:E2')
            self.assertFalse(book['Listings'].protection.sheet)
            self.assertEqual(book['Listings']['A1'].fill.fgColor.rgb,'00EEEEEE')

    def test_never_silently_truncates_collected_data(self):
        data=self.data();data['tables'][0]['rows'][0][3]='x'*32768
        with TemporaryDirectory() as temp:
            with self.assertRaisesRegex(ValueError,'refusing to truncate'):
                create(data,Path(temp)/'listings.xlsx')


if __name__ == '__main__':unittest.main()
