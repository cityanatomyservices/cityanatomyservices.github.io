"""Create a native-Sheets-ready workbook; all source strings are literal text."""
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ZONE = ZoneInfo('America/Chicago')
WORKFLOW = 'https://github.com/cityanatomyservices/cityanatomyservices.github.io/actions/workflows/export-explore.yml'


def local(value):
    return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(ZONE).strftime('%b %d, %Y %I:%M %p %Z')


def put(sheet, row, values):
    for col, value in enumerate(values, 1):
        if isinstance(value, str) and len(value) > 32767:
            raise ValueError('A source field exceeds the spreadsheet cell limit; refusing to truncate it.')
        cell = sheet.cell(row, col, value)
        if isinstance(value, str):
            cell.data_type = 's'  # Organizer text beginning with = must never become a formula.
        cell.alignment = Alignment(vertical='top', wrap_text=True)


def create(data, destination, links=None):
    links = links or {}
    book = Workbook()
    sheet = book.active
    sheet.title = 'Start here'
    summary = [
        ['Austin.chat Explore export', data['title']],
        ['Created', local(data['exported'])],
        ['Selection', data['when'].title()],
        ['Time window', local(data['windowStart'])+' through '+local(data['windowEnd'])+' (end exclusive)'],
        ['Listings / map places', f"{len(data['records'])} listings / {len(data['venues'])} numbered places"],
        ['Scope', 'All six Explore categories, all locations, all theaters, Free & cheap up to $20. No search or nearby restriction.'],
        ['Research columns', 'Use columns E–H in Listings for your research, extra sources, Substack draft and editorial status.'],
        ['Snapshots', 'Every export creates a new spreadsheet. Later exports never refresh or overwrite this file or your notes.'],
        ['Map image', links.get('Map.png', 'Map.png in the same export folder')],
        ['Interactive map', links.get('Map.html', 'Map.html in the same export folder')],
        ['Opening the map', 'Download Map.html from Drive and open it in your browser. The listings are embedded; internet access loads the basemap. Click numbered pins or category buttons.'],
        ['GIS data', links.get('Listings.geojson', 'Listings.geojson in the same export folder')],
        ['Map numbers', 'Map pin is a place number shared by Listings, Movie showtimes and Map places. Multiple listings at one location share a pin.'],
        ['Movie showtimes', 'Movie showtimes lists every matching start and ticket URL. Collected data also preserves the complete original movie record, including other showtimes.'],
        ['Overlapping categories', 'Listings follows Explore’s deduplication. All categories identifies other matching categories; Collected data retains every original variant and price.'],
        ['Unknown end times', 'Ends (Austin) is blank when no end was published. Raw End ISO may be a technical expiry boundary, not an actual finish time.'],
        ['Sources and completeness', 'Sources records every feed/source status. Missing category files stop export. Stale or partly unavailable sources are flagged, never filled with invented information.'],
        ['Source notes', '\n'.join(data['warnings']) or 'No delayed or unavailable sources reported in these snapshots.'],
        ['Time rules', 'Same as Explore Austin: Tonight is 6 PM–3 AM; Weekend is Saturday–Monday midnight. Before 3 AM, Tonight refers to the preceding evening. Already-ended/expired listings are omitted.'],
        ['Generate another export', WORKFLOW],
        ['How to run', 'Sign into GitHub as the site owner → Run workflow → choose Tonight or Weekend → Run workflow. Wait for success, then open the new dated folder in Drive.'],
        ['Run reference', data['runId']],
        ['Privacy', 'Files are private in Drive. No public export button, public download, automatic sharing or Substack posting is added.'],
        ['Attribution', 'Map: OpenFreeMap / © OpenMapTiles / © OpenStreetMap contributors. Event information: organizer links in Listings and Sources.']
    ]
    for i, row in enumerate(summary, 1): put(sheet, i, row)
    sheet.column_dimensions['A'].width = 28
    sheet.column_dimensions['B'].width = 105
    for row in range(1, len(summary)+1): sheet.row_dimensions[row].height = 44 if row > 4 else 28
    sheet.freeze_panes = 'B2'
    sheet['A1'].font = sheet['B1'].font = Font(bold=True, size=12)
    for table in data['tables']:
        sheet = book.create_sheet(table['name'])
        headers = table['headers']
        put(sheet, 1, headers)
        for i, values in enumerate(table['rows'], 2):
            if len(values) != len(headers): raise ValueError('Spreadsheet row/header mismatch: '+table['name'])
            put(sheet, i, values)
            sheet.row_dimensions[i].height = 60 if not table.get('technical') else 32
        sheet.freeze_panes = 'E2' if table['name'] == 'Listings' else 'C2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(1,len(table["rows"])+1)}'
        sheet.row_dimensions[1].height = 38
        for i, header in enumerate(headers, 1):
            width = 43 if header in table.get('wide', []) else 23
            if header in ('Listing #', 'Map pin'): width = 11
            if 'JSON' in header: width = 50
            sheet.column_dimensions[get_column_letter(i)].width = width
            cell = sheet.cell(1, i)
            cell.fill = PatternFill('solid', fgColor='EEEEEE')
            cell.font = Font(bold=True)
            if header in table.get('notes', []):
                from openpyxl.comments import Comment
                cell.comment = Comment('Your editable research column. Later exports will not overwrite it.', 'Austin.chat')
        sheet.sheet_view.zoomScale = 90
    book.save(Path(destination))
